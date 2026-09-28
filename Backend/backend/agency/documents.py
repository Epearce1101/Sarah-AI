"""Real documents: Word (.docx), Excel (.xlsx), PDF, CSV, HTML, Markdown, text.

Content is written in simple markdown ("# Heading", "- bullet", "1. item",
"**bold**", blank line = new paragraph); tables are rows (lists). Every save
is read back from disk and the result says what's actually in the file, so
"saved" is a fact, not a hope.
"""
from __future__ import annotations

import csv
import html
import io
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

FORMATS = {".txt", ".md", ".docx", ".xlsx", ".csv", ".pdf", ".html", ".htm"}


def _blocks(content: str) -> List[tuple]:
    """markdown-ish text -> [(kind, text)]: h1/h2/h3, bullet, number, para, pagebreak."""
    out: List[tuple] = []
    para: List[str] = []

    def flush():
        if para:
            out.append(("para", " ".join(s.strip() for s in para)))
            para.clear()

    for raw in (content or "").replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            flush()
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            flush()
            out.append((f"h{len(m.group(1))}", m.group(2).strip()))
        elif re.match(r"^\s*[-*•]\s+", line):
            flush()
            out.append(("bullet", re.sub(r"^\s*[-*•]\s+", "", line)))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            flush()
            out.append(("number", re.sub(r"^\s*\d+[.)]\s+", "", line)))
        elif line.strip() in ("---", "***", "<pagebreak>"):
            flush()
            out.append(("pagebreak", ""))
        else:
            para.append(line)
    flush()
    return out


def _runs(text: str) -> List[tuple]:
    """'a **b** c' -> [('a ', False), ('b', True), (' c', False)]"""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    return [(p[2:-2], True) if p.startswith("**") and p.endswith("**") else (p, False) for p in parts if p]


def _plain(text: str) -> str:
    return text.replace("**", "")


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------

def _write_docx(path: Path, title: str, content: str, append: bool) -> None:
    import docx

    doc = docx.Document(str(path)) if append and path.exists() else docx.Document()
    if title and not append:
        doc.add_heading(title, level=0)
    for kind, text in _blocks(content):
        if kind.startswith("h"):
            doc.add_heading(_plain(text), level=int(kind[1]))
        elif kind == "pagebreak":
            doc.add_page_break()
        else:
            style = {"bullet": "List Bullet", "number": "List Number"}.get(kind)
            p = doc.add_paragraph(style=style)
            for chunk, bold in _runs(text):
                p.add_run(chunk).bold = bold
    doc.save(str(path))


def _rows_from(content: str, rows: Optional[List[List[Any]]]) -> List[List[Any]]:
    if rows:
        return [list(r) if isinstance(r, (list, tuple)) else [r] for r in rows]
    lines = [l for l in (content or "").splitlines() if l.strip()]
    if not lines:
        return []
    delim = "\t" if "\t" in lines[0] else ("|" if lines[0].count("|") >= 2 else ",")
    out = []
    for l in lines:
        if delim == "|":
            if re.fullmatch(r"\|?[\s:|-]+\|?", l):
                continue  # markdown table separator
            out.append([c.strip() for c in l.strip().strip("|").split("|")])
        else:
            out.append(next(csv.reader([l], delimiter=delim)))
    return out


def _cell(v: Any) -> Any:
    if isinstance(v, str):
        s = v.strip()
        if re.fullmatch(r"-?\d+", s):
            return int(s)
        if re.fullmatch(r"-?\d+\.\d+", s):
            return float(s)
    return v


def _write_xlsx(path: Path, title: str, rows: List[List[Any]], sheet: str, append: bool) -> None:
    import openpyxl
    from openpyxl.styles import Font

    if append and path.exists():
        wb = openpyxl.load_workbook(str(path))
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        start = ws.max_row + 1
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = (sheet or title or "Sheet1")[:31]
        start = 1
    for r in rows:
        ws.append([_cell(v) for v in r])
    if start == 1 and rows:
        for c in ws[1]:
            c.font = Font(bold=True)
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(60, max(8, width + 2))
    wb.save(str(path))


def _write_csv(path: Path, rows: List[List[Any]], append: bool) -> None:
    with path.open("a" if append else "w", newline="", encoding="utf-8-sig" if not append else "utf-8") as fh:
        csv.writer(fh).writerows(rows)


def _font_files() -> Optional[tuple]:
    """(regular, bold) Windows fonts with wide Unicode coverage."""
    fonts = Path(r"C:\Windows\Fonts")
    for regular, bold in (("segoeui.ttf", "segoeuib.ttf"), ("arial.ttf", "arialbd.ttf"), ("calibri.ttf", "calibrib.ttf")):
        if (fonts / regular).exists():
            b = fonts / bold
            return str(fonts / regular), str(b if b.exists() else fonts / regular)
    return None


def _write_pdf(path: Path, title: str, content: str) -> None:
    from fpdf import FPDF

    from fpdf.enums import XPos, YPos

    nl = {"new_x": XPos.LMARGIN, "new_y": YPos.NEXT}  # each block starts on a new line
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    font = _font_files()
    if font:
        pdf.add_font("Body", "", font[0])
        pdf.add_font("Body", "B", font[1])
        family = "Body"
    else:
        family = "Helvetica"
    strip = (lambda s: s) if font else (lambda s: s.encode("latin-1", "replace").decode("latin-1"))
    no_emoji = lambda s: re.sub("[\U00010000-\U0010FFFF]", "", s)
    if title:
        pdf.set_font(family, "B", 20)
        pdf.multi_cell(0, 10, strip(no_emoji(title)), **nl)
        pdf.ln(3)
    for kind, text in _blocks(content):
        text = strip(no_emoji(_plain(text)))
        if kind.startswith("h"):
            pdf.set_font(family, "B", {"h1": 16, "h2": 14, "h3": 12}[kind])
            pdf.ln(2)
            pdf.multi_cell(0, 8, text, **nl)
        elif kind == "pagebreak":
            pdf.add_page()
        else:
            pdf.set_font(family, "", 11)
            prefix = "•  " if kind == "bullet" and font else ("-  " if kind == "bullet" else "")
            pdf.multi_cell(0, 6, prefix + text, **nl)
            pdf.ln(1 if kind in ("bullet", "number") else 3)
    pdf.output(str(path))


def _write_html(path: Path, title: str, content: str) -> None:
    body, in_list = [], None
    for kind, text in _blocks(content):
        tag = {"bullet": "ul", "number": "ol"}.get(kind)
        if in_list and tag != in_list:
            body.append(f"</{in_list}>")
            in_list = None
        inline = "".join(f"<b>{html.escape(c)}</b>" if b else html.escape(c) for c, b in _runs(text))
        if tag:
            if not in_list:
                body.append(f"<{tag}>")
                in_list = tag
            body.append(f"<li>{inline}</li>")
        elif kind.startswith("h"):
            body.append(f"<{kind}>{inline}</{kind}>")
        elif kind == "pagebreak":
            body.append("<hr>")
        else:
            body.append(f"<p>{inline}</p>")
    if in_list:
        body.append(f"</{in_list}>")
    head = f"<title>{html.escape(title)}</title>" if title else ""
    h1 = f"<h1>{html.escape(title)}</h1>" if title else ""
    path.write_text(f"<!DOCTYPE html><html><head><meta charset='utf-8'>{head}</head><body>{h1}{''.join(body)}</body></html>",
                    encoding="utf-8")


# ---------------------------------------------------------------------------
# Reading (also the proof after every save)
# ---------------------------------------------------------------------------

def read(path: Path, max_chars: int = 8000) -> Dict[str, Any]:
    ext = path.suffix.lower()
    if not path.exists():
        raise FileNotFoundError(f"{path} doesn't exist")
    info: Dict[str, Any] = {"path": str(path), "bytes": path.stat().st_size}
    if ext == ".docx":
        import docx
        doc = docx.Document(str(path))
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        info.update(paragraphs=len(paras), text="\n".join(paras)[:max_chars])
    elif ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        sheets = {}
        for ws in wb.worksheets:
            rows = [[("" if v is None else v) for v in r] for r in ws.iter_rows(values_only=True)]
            sheets[ws.title] = {"rows": len(rows), "first_rows": rows[:15]}
        info["sheets"] = sheets
    elif ext == ".pdf":
        data = path.read_bytes()
        info.update(pages=len(re.findall(rb"/Type\s*/Page[^s]", data)), valid_pdf=data.startswith(b"%PDF"))
    else:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        info.update(lines=text.count("\n") + 1, text=text[:max_chars])
    return info


def open_for_zero(path: Path) -> str:
    """Open the file in whatever can show it on this PC (WordPad for Word files
    when Word isn't installed)."""
    import os
    import subprocess

    try:
        os.startfile(str(path))
        return "opened"
    except OSError:
        pass
    wordpad = Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "Windows NT" / "Accessories" / "wordpad.exe"
    if path.suffix.lower() in (".docx", ".txt", ".md", ".csv", ".rtf") and wordpad.exists():
        subprocess.Popen([str(wordpad), str(path)])
        return "opened in WordPad (no Word on this PC)"
    if path.suffix.lower() in (".html", ".htm", ".pdf"):
        from .desktop import chrome_path
        chrome = chrome_path()
        if chrome:
            subprocess.Popen([chrome, str(path)])
            return "opened in Chrome"
    return f"saved, but nothing on this PC opens {path.suffix} files (no Excel/Office installed)"


def write(path: Path, content: str = "", title: str = "", rows: Optional[List[List[Any]]] = None,
          sheet: str = "", append: bool = False) -> Dict[str, Any]:
    ext = path.suffix.lower()
    if ext not in FORMATS:
        raise ValueError(f"can't make {ext or 'extension-less'} files; use one of {', '.join(sorted(FORMATS))}")
    path.parent.mkdir(parents=True, exist_ok=True)
    if ext == ".docx":
        _write_docx(path, title, content, append)
    elif ext == ".xlsx":
        data = _rows_from(content, rows)
        if not data:
            raise ValueError("an .xlsx needs rows (a list of lists) or table-like content")
        _write_xlsx(path, title, data, sheet, append)
    elif ext == ".csv":
        data = _rows_from(content, rows)
        if not data:
            raise ValueError("a .csv needs rows or table-like content")
        _write_csv(path, data, append)
    elif ext == ".pdf":
        if append:
            raise ValueError("PDFs can't be appended to; write the whole document again")
        _write_pdf(path, title, content)
    elif ext in (".html", ".htm"):
        _write_html(path, title, content)
    else:
        text = (f"{title}\n\n" if title and not append else "") + (content or "")
        lead = "\n" if append and path.exists() and path.stat().st_size else ""
        with path.open("a" if append else "w", encoding="utf-8") as fh:
            fh.write(lead + text)
    proof = read(path, 1500)
    proof["check"] = "read back from disk after saving"
    return proof
