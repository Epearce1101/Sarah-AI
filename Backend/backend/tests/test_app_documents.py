"""Documents, Zero's real folders, and checking her work before she answers."""
import asyncio
import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from backend.agency import desktop, documents, tools
from backend.tests.test_agency import FakeStream, _chunk, _tc


def run(coro):
    return asyncio.run(coro)


LETTER = "# Trip plan\n\nWe leave **Friday** at 9.\n\n## Packing\n- passport\n- charger\n\n1. book hotel\n2. pack"


def test_word_document_is_real_and_read_back(tmp_path):
    out = documents.write(tmp_path / "trip.docx", content=LETTER, title="Osaka")
    assert out["check"] and out["paragraphs"] >= 7
    assert "Osaka" in out["text"] and "passport" in out["text"] and "**" not in out["text"]
    import docx
    doc = docx.Document(str(tmp_path / "trip.docx"))
    styles = [p.style.name for p in doc.paragraphs]
    assert "Title" in styles and "Heading 1" in styles and "List Bullet" in styles and "List Number" in styles
    assert any(r.bold for p in doc.paragraphs for r in p.runs if r.text == "Friday")
    more = documents.write(tmp_path / "trip.docx", content="- sunscreen", append=True)
    assert "sunscreen" in more["text"] and "passport" in more["text"]


def test_spreadsheet_from_rows_and_from_a_markdown_table(tmp_path):
    out = documents.write(tmp_path / "budget.xlsx", rows=[["Item", "Cost"], ["Hotel", "320"], ["Food", 95.5]])
    sheet = next(iter(out["sheets"].values()))
    assert sheet["rows"] == 3 and sheet["first_rows"][1] == ["Hotel", 320]   # numbers are numbers
    out = documents.write(tmp_path / "t.xlsx", content="| Game | Hours |\n|---|---|\n| Elden Ring | 120 |")
    assert next(iter(out["sheets"].values()))["first_rows"] == [["Game", "Hours"], ["Elden Ring", 120]]
    with pytest.raises(ValueError):
        documents.write(tmp_path / "empty.xlsx", content="")


def test_pdf_csv_html_and_text(tmp_path):
    pdf = documents.write(tmp_path / "note.pdf", content=LETTER + "\n\nÜnïcödé ✓ 🎉", title="Plan")
    assert pdf["valid_pdf"] and pdf["pages"] >= 1 and pdf["bytes"] > 500
    csv_out = documents.write(tmp_path / "list.csv", rows=[["a", "b"], [1, 2]])
    assert "a,b" in csv_out["text"]
    html_out = documents.write(tmp_path / "page.html", content=LETTER, title="T")
    assert "<h1>Trip plan</h1>" in html_out["text"] and "<li>passport</li>" in html_out["text"]
    txt = documents.write(tmp_path / "a.txt", content="hello", title="Hi")
    assert txt["text"] == "Hi\n\nhello"
    assert documents.write(tmp_path / "a.txt", content="more", append=True)["text"] == "Hi\n\nhello\nmore"
    with pytest.raises(ValueError):
        documents.write(tmp_path / "x.exe", content="nope")


def test_desktop_and_documents_are_zeros_real_folders(tmp_path, monkeypatch):
    monkeypatch.setattr(desktop, "known_folder", lambda name: {"desktop": tmp_path / "Desk",
                                                               "my documents": tmp_path / "Docs"}.get(name.lower()))
    assert tools._resolve("Desktop/letter.txt") == tmp_path / "Desk" / "letter.txt"
    assert tools._resolve("My Documents/a.docx") == tmp_path / "Docs" / "a.docx"
    assert tools._resolve("notes/x.md").parts[-3:] == ("sarah_workspace", "notes", "x.md")


def test_document_tool_saves_where_asked(tmp_path, monkeypatch):
    monkeypatch.setattr(desktop, "known_folder", lambda name: tmp_path if name.lower() == "desktop" else None)
    out = run(tools.call("document", {"action": "create", "path": "Desktop/Shopping.docx", "content": "- milk\n- eggs"}))
    assert out["ok"] and (tmp_path / "Shopping.docx").exists() and "milk" in out["result"]


def test_what_counts_as_checking():
    assert tools.is_check("app", {"action": "read"}) and tools.is_check("read_file", {})
    assert tools.is_check("browser", {"action": "read"}) and not tools.is_check("browser", {"action": "click"})
    assert tools.unverified_action("control_input", {"action": "type"}, True, "Typed") == "used the keyboard (type)"
    assert tools.unverified_action("app", {"action": "keys", "window": "Notepad"}, True, "") == "keys in Notepad"
    assert tools.unverified_action("app", {"action": "type"}, True, "") is None     # it reads the field back
    assert tools.unverified_action("browser", {"action": "click"}, True, "Nothing on the page changed") is not None
    assert tools.unverified_action("browser", {"action": "click"}, True, '{"changed": true}') is None


def test_she_must_look_before_claiming_it_worked(monkeypatch):
    from backend.memory import openrouter_client as oc

    seen, n = [], [0]

    class Completions:
        def create(self, **kw):
            n[0] += 1
            seen.append(kw["messages"])
            if n[0] == 1:
                args = json.dumps({"action": "screen_size"})  # harmless stand-in; patched below as a keyboard action
                return FakeStream([_chunk(tool_calls=[_tc(0, "k", "control_input", args)], finish="tool_calls")])
            if n[0] == 2:
                return FakeStream([_chunk("Typed it all, done!", finish="stop")])
            if n[0] == 3:
                return FakeStream([_chunk(tool_calls=[_tc(0, "r", "list_my_tools", "{}")], finish="tool_calls")])
            return FakeStream([_chunk("Checked: it's there.", finish="stop")])

    client = oc.OpenRouterClient.__new__(oc.OpenRouterClient)
    client._client = SimpleNamespace(chat=SimpleNamespace(completions=Completions()))
    client.config = SimpleNamespace(llm_max_completion_tokens=100, llm_temperature=0.5, chars_per_token=4)
    client._is_local_mode = lambda: False
    packet = SimpleNamespace(messages=[{"role": "user", "content": "type hi"}], estimated_tokens=10, debug_info={})
    client._prepare_turn = lambda *a, **k: (packet, 1)
    client._finish_turn = lambda **kw: SimpleNamespace(content=kw["raw_content"])
    monkeypatch.setattr(oc.llm_models, "completion_kwargs", lambda: {"model": "fake:free", "extra_body": None})
    monkeypatch.setattr(tools, "_custom_tools", lambda: {})
    monkeypatch.setattr(tools, "unverified_action", lambda name, args, ok, result: "used the keyboard (type)"
                        if name == "control_input" else None)

    async def collect():
        return [e async for e in client.chat_stream(1, "type hi")]

    events = run(collect())
    assert "haven't checked the result" in seen[2][-1]["content"]
    assert events[-1]["response"].content.endswith("Checked: it's there.")
    assert n[0] == 4  # one nudge only


def test_close_and_save_on_a_new_file_says_not_saved(monkeypatch):
    wins = [{"hwnd": 1, "title": "Untitled - Notepad", "pid": 7, "app": "notepad.exe"}]
    monkeypatch.setattr(desktop, "windows", lambda: wins)
    monkeypatch.setattr(desktop, "_alive", lambda h: True)
    monkeypatch.setattr(desktop, "require_unlocked", lambda: None)
    monkeypatch.setattr(desktop, "focus", lambda t: wins[0])
    monkeypatch.setattr(desktop.user32, "PostMessageW", lambda *a: None)
    import pyautogui

    def press_save(*keys):
        wins.append({"hwnd": 2, "title": "Save As", "pid": 7, "app": "notepad.exe"})

    monkeypatch.setattr(pyautogui, "hotkey", press_save)
    monkeypatch.setattr(pyautogui, "press", lambda k: None)
    with pytest.raises(desktop.DesktopError, match="NOT saved"):
        desktop.close("notepad", save=True, wait=0.3)
